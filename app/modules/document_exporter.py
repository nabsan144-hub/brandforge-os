"""Local Unicode PDF/DOCX export. PDFs embed OFL fonts and use HarfBuzz/bidi.

Nothing is fetched during export. Documents treat campaign content as literal
text, never HTML/XML instructions. Oversized content is rejected, not cut.
"""
import base64
import io
import re
from pathlib import Path
from typing import Any, Dict

FONTS = Path(__file__).resolve().parents[1] / 'brandforge_assets' / 'fonts'


def _text(value: Any, limit: int = 30000) -> str:
    value = str(value or '')
    if len(value) > limit:
        raise ValueError(f'Document section exceeds {limit:,} characters; split it before export')
    return ''.join(ch for ch in value if ch in '\n\r\t' or ord(ch) >= 32)


def _pdf_safe(value: Any, limit: int = 30000) -> str:
    # Backward-compatible helper; the PDF renderer itself does not parse markup.
    import html
    return html.escape(_text(value, limit), quote=False).replace('\n', '<br/>')


def _parts(campaign):
    strategy, copy, analysis = (campaign.get(k) or {} for k in ('strategy','copy','analysis'))
    labels={'ur':['حکمت عملی','حتمی مسودہ','SEO اور تبادلوں کی تجاویز'],'hi':['रणनीति','अंतिम मसौदा','SEO और रूपांतरण सुझाव'],'es':['Estrategia','Texto final','Sugerencias SEO y CRO'],'pt':['Estratégia','Texto final','Sugestões de SEO e CRO']}.get(campaign.get('lang'),['Strategy','Final copy','SEO and conversion suggestions'])
    return strategy,list(zip(labels,[strategy.get('strategy_text'),copy.get('copy_text'),analysis.get('seo_analysis')]))


def _display(line):
    line = re.sub(r'^#{1,6}\s+', '', line)
    line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
    return line[2:] if line.startswith('> ') else line


def _logo(strategy):
    value = str(strategy.get('approved_logo') or '')
    if re.fullmatch(r'data:image/png;base64,[A-Za-z0-9+/]+={0,2}', value) and len(value) <= 500000:
        return io.BytesIO(base64.b64decode(value.split(',')[1]))
    return None


def campaign_docx(campaign: Dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'Noto Sans'
    normal.font.size = Pt(11)
    strategy, parts = _parts(campaign)
    logo = _logo(strategy)
    if logo: document.add_picture(logo, width=Inches(1.1))
    document.add_heading(_text(strategy.get('product_name'), 80) or 'Campaign', 0)
    document.add_paragraph(f"Campaign: {_text(campaign.get('campaign_name'),80)}")
    document.add_paragraph(f"Status: {_text(campaign.get('status'),40).replace('_',' ').title()}")
    for title, text in parts:
        document.add_heading(title, level=1)
        for line in _text(text).split('\n'):
            p = document.add_paragraph()
            rtl = bool(re.match(r'^[^A-Za-z\u0900-\u097f]*[\u0600-\u06ff]', line))
            if rtl:
                bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); p._p.get_or_add_pPr().append(bidi)
            run = p.add_run(_display(line))
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn('w:cs'), 'Noto Naskh Arabic' if re.search(r'[\u0600-\u06ff]',line) else 'Noto Sans Devanagari')
    document.add_heading({'ur':'اشاعت سے پہلے جائزہ','hi':'प्रकाशन से पहले समीक्षा','es':'Revisión','pt':'Revisão'}.get(campaign.get('lang'),'Review note'), 1)
    document.add_paragraph('Verify facts, offers, destinations, legal requirements, platform policies and client approval before publishing. Fonts for supported scripts are included in the Desktop installation; install them for consistent editable DOCX layout. The PDF embeds its fonts.')
    stream = io.BytesIO(); document.save(stream); return stream.getvalue()


def campaign_pdf(campaign: Dict[str, Any]) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos
    from fontTools.ttLib import TTFont
    pdf = FPDF(format='A4')
    pdf.set_margins(17,16,17); pdf.set_auto_page_break(True,17)
    names = [('Brand','NotoSans'),('Arabic','NotoNaskhArabic'),('Hindi','NotoSansDevanagari')]
    coverage = set()
    for family, filename in names:
        path = FONTS / f'{filename}-Regular.ttf'
        pdf.add_font(family, fname=str(path))
        with TTFont(path) as font: coverage.update(font.getBestCmap())
    pdf.set_fallback_fonts(['Arabic','Hindi'])
    pdf.set_text_shaping(True)
    pdf.set_title(_text(campaign.get('campaign_name'),80) or 'Campaign')
    pdf.set_author('BrandForge OS')
    pdf.set_lang(campaign.get('lang') or 'en')
    pdf.add_page()
    strategy, parts = _parts(campaign)
    logo = _logo(strategy)
    if logo: pdf.image(logo,x=17,y=pdf.y,w=24,h=24,keep_aspect_ratio=True); pdf.ln(27)

    def paragraph(value, size=10.5, spacing=5.8):
        pdf.set_font('Brand', size=size)
        value = _text(value)
        # Remove only unsupported decorative symbols/emoji. Unsupported letter
        # scripts get an explicit error instead of invisible glyph boxes.
        import unicodedata
        safe = ''
        for ch in value:
            if ch in '\n\t\r' or ord(ch) in coverage: safe += ch
            elif unicodedata.category(ch).startswith(('L','N')):
                raise ValueError('This document contains an unsupported script; use the text/HTML export or add a licensed font')
        for line in safe.split('\n'):
            if not line.strip(): pdf.ln(2.2); continue
            pdf.multi_cell(0,spacing,_display(line),align='R' if campaign.get('lang')=='ur' else 'L',new_x=XPos.LMARGIN,new_y=YPos.NEXT)

    paragraph(strategy.get('product_name') or 'Campaign',22,11)
    paragraph(f"Campaign: {campaign.get('campaign_name','')}")
    paragraph(f"Status: {campaign.get('status','draft').replace('_',' ').title()}")
    for title, text in parts:
        pdf.ln(6); paragraph(title,15,8); paragraph(text)
    pdf.ln(6); paragraph({'ur':'اشاعت سے پہلے جائزہ','hi':'प्रकाशन से पहले समीक्षा','es':'Revisión','pt':'Revisão'}.get(campaign.get('lang'),'Review note'),15,8)
    paragraph('Verify factual claims, legal requirements, platform policies, destination links and client approval before publishing. This is a draft, not an approval or performance guarantee.')
    return bytes(pdf.output())
